from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import cgi
import json
import os
import time
import unicodedata

from openpyxl import Workbook, load_workbook


APP_DIR = Path(os.environ.get("APP_DIR", "/app"))
DATA_DIR = Path(os.environ.get("DATA_DIR", "/data"))
HTML_FILE = APP_DIR / "index.html"
DATA_FILE = DATA_DIR / "tickets.json"
MAX_UPLOAD_BYTES = int(os.environ.get("MAX_UPLOAD_BYTES", str(10 * 1024 * 1024)))

COLUMNS = [
    ("item", "Item"),
    ("ticket", "Ticket"),
    ("aplicacion", "Aplicación"),
    ("detalle", "Detalle"),
    ("fecha_estimada", "Fecha estimada de resolución"),
    ("horas_estimadas", "Horas estimadas desarrollo"),
    ("estado", "Estado"),
]

HEADER_ALIASES = {
    "item": {"item"},
    "ticket": {"ticket", "tickets", "n ticket", "n° ticket"},
    "aplicacion": {"aplicacion", "aplicación"},
    "detalle": {"detalle", "descripcion", "descripción"},
    "fecha_estimada": {
        "fecha estimada de resolucion",
        "fecha estimada de resolución",
        "fecha estimada",
        "fecha",
    },
    "horas_estimadas": {
        "horas estimadas desarrollo",
        "horas estimadas de desarrollo",
        "horas estimadas",
        "horas",
    },
    "estado": {"estado", "status"},
}

STATIC_FILES = {
    "/app.js": ("app.js", "application/javascript; charset=utf-8"),
    "/icon.svg": ("icon.svg", "image/svg+xml"),
}


def normalize_header(value) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.lower().strip().split())


def match_field(header: str) -> str | None:
    normalized = normalize_header(header)
    for field, aliases in HEADER_ALIASES.items():
        if normalized in aliases:
            return field
    return None


def format_cell_value(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def empty_row() -> dict:
    return {field: "" for field, _ in COLUMNS}


def default_row(**overrides) -> dict:
    row = empty_row()
    row["estado"] = "abierto"
    row.update(overrides)
    return row


def normalize_estado(value: str) -> str:
    text = normalize_header(value)
    if text in {"cerrado", "closed", "resuelto", "finalizado"}:
        return "cerrado"
    return "abierto"


def parse_workbook(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row_idx = None
    mapping: dict[int, str] = {}

    for idx, row in enumerate(rows):
        current = {}
        for col_idx, cell in enumerate(row):
            field = match_field(cell)
            if field:
                current[col_idx] = field
        if len(current) >= 3:
            header_row_idx = idx
            mapping = current
            break

    if header_row_idx is None:
        return []

    tickets: list[dict] = []
    for row in rows[header_row_idx + 1 :]:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        ticket = default_row()
        for col_idx, field in mapping.items():
            if col_idx < len(row):
                value = format_cell_value(row[col_idx])
                if field == "estado":
                    ticket[field] = normalize_estado(value) if value else "abierto"
                else:
                    ticket[field] = value
        if not any(ticket[field] for field in ("item", "ticket", "aplicacion", "detalle")):
            continue
        tickets.append(ticket)
    return tickets


def load_tickets() -> list[dict]:
    if not DATA_FILE.exists():
        seed = APP_DIR / "seed.xlsx"
        if seed.exists():
            tickets = parse_workbook(seed)
            if tickets:
                save_tickets(tickets)
                return tickets
        return []
    payload = json.loads(DATA_FILE.read_text(encoding="utf-8"))
    tickets = payload.get("tickets", [])
    return [default_row(**ticket) for ticket in tickets]


def save_tickets(tickets: list[dict]) -> dict:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "updatedAt": int(time.time()),
        "tickets": tickets,
    }
    DATA_FILE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def export_workbook(tickets: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    ws.append([label for _, label in COLUMNS])
    for ticket in tickets:
        ws.append([ticket.get(field, "") for field, _ in COLUMNS])
    wb.save(path)


class Handler(BaseHTTPRequestHandler):
    server_version = "SustentableTickets/1.0"

    def log_message(self, fmt, *args):
        return

    def send_json(self, status: int, payload: dict | list):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_file(self, path: Path, content_type: str):
        body = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path

        if path in ("", "/"):
            return self.send_file(HTML_FILE, "text/html; charset=utf-8")

        if path in STATIC_FILES:
            filename, content_type = STATIC_FILES[path]
            file_path = APP_DIR / filename
            if file_path.exists():
                return self.send_file(file_path, content_type)

        if path == "/api/tickets":
            payload = load_tickets()
            meta = {}
            if DATA_FILE.exists():
                meta = json.loads(DATA_FILE.read_text(encoding="utf-8"))
            return self.send_json(200, {"tickets": payload, "updatedAt": meta.get("updatedAt")})

        if path == "/api/export":
            tickets = load_tickets()
            export_path = DATA_DIR / "export.xlsx"
            DATA_DIR.mkdir(parents=True, exist_ok=True)
            export_workbook(tickets, export_path)
            body = export_path.read_bytes()
            self.send_response(200)
            self.send_header(
                "Content-Type",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.send_header(
                "Content-Disposition",
                'attachment; filename="sustentable-tickets.xlsx"',
            )
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        self.send_error(404)

    def do_POST(self):
        parsed = urlparse(self.path)
        path = parsed.path
        length = int(self.headers.get("Content-Length", "0"))

        if path == "/api/tickets":
            body = self.rfile.read(length)
            try:
                payload = json.loads(body.decode("utf-8"))
                tickets = payload.get("tickets", [])
                normalized = []
                for ticket in tickets:
                    row = default_row()
                    for field, _ in COLUMNS:
                        if field in ticket:
                            value = str(ticket.get(field, "") or "").strip()
                            if field == "estado":
                                row[field] = normalize_estado(value)
                            else:
                                row[field] = value
                    normalized.append(row)
                saved = save_tickets(normalized)
                return self.send_json(200, saved)
            except (json.JSONDecodeError, TypeError):
                return self.send_json(400, {"error": "JSON inválido"})

        if path == "/api/import":
            form = cgi.FieldStorage(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                    "CONTENT_LENGTH": str(length),
                },
            )
            file_item = form["file"] if "file" in form else None
            if not file_item or not getattr(file_item, "filename", None):
                return self.send_json(400, {"error": "No se recibió ningún archivo"})
            data = file_item.file.read(MAX_UPLOAD_BYTES + 1)
            if len(data) > MAX_UPLOAD_BYTES:
                return self.send_json(400, {"error": "Archivo demasiado grande"})
            upload_path = DATA_DIR / "upload.xlsx"
            DATA_DIR.mkdir(parents=True, exist_ok=True)
            upload_path.write_bytes(data)
            tickets = parse_workbook(upload_path)
            if not tickets:
                return self.send_json(400, {"error": "No se pudieron leer filas del Excel"})
            saved = save_tickets(tickets)
            return self.send_json(200, saved)

        self.send_error(404)


def main():
    APP_DIR.mkdir(parents=True, exist_ok=True)
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    port = int(os.environ.get("PORT", "8000"))
    server = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    print(f"Sustentable tickets escuchando en :{port}")
    server.serve_forever()


if __name__ == "__main__":
    main()
